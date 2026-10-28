"""Bounded file loading. Never silently repair malformed rows or cast CSV cells."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

import polars as pl

MAX_CELLS = 5_000_000


def _row_budget(headers, requested):
    _validate_headers(headers)
    if len(headers) > 500:
        raise ValueError("MVP поддерживает до 500 столбцов.")
    return min(requested, max(1, MAX_CELLS // len(headers)))


@dataclass
class Dataset:
    frame: pl.DataFrame
    source: str
    format: str
    truncated: bool
    notes: list[str]


def excel_sheets(path: str | Path) -> list[str]:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def load_dataset(path: str | Path, *, max_rows: int = 200_000,
                 separator: str | None = None, encoding: str = "utf-8-sig",
                 sheet: str | None = None) -> Dataset:
    path = Path(path)
    if max_rows < 1:
        raise ValueError("Лимит строк должен быть положительным.")
    if not path.is_file():
        raise ValueError(f"Файл не найден: {path}")
    if path.stat().st_size > 256 * 1024 * 1024:
        raise ValueError("MVP поддерживает файлы до 256 МБ. Подготовьте меньший файл.")
    ext = path.suffix.lower()
    notes: list[str] = []
    if ext == ".csv":
        # Python's strict reader supports cp1251, BOM, quoted newlines and row limits.
        with path.open(encoding=encoding, newline="") as stream:
            sample = stream.read(65536)
            stream.seek(0)
            if not separator:
                try:
                    separator = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    # A malformed body can defeat Sniffer; use header width so
                    # the strict row validator still exposes the broken record.
                    widths = {}
                    for candidate in (",", ";", "\t", "|"):
                        widths[candidate] = len(next(csv.reader(io.StringIO(sample), delimiter=candidate), []))
                    separator = max(widths, key=widths.get)
            if len(separator) != 1:
                raise ValueError("Разделитель должен состоять из одного символа.")
            reader = csv.reader(stream, delimiter=separator, strict=True)
            try:
                headers = next(reader)
            except StopIteration:
                raise ValueError("Файл пуст: нет заголовка.") from None
            budget = _row_budget(headers, max_rows)
            rows = []
            for row in reader:
                if not row:  # Explicitly skip empty physical records.
                    continue
                if len(row) != len(headers):
                    raise ValueError(f"CSV, строка {reader.line_num}: ожидалось {len(headers)} полей, получено {len(row)}.")
                rows.append(row)
                if len(rows) > budget:
                    break
        frame = pl.DataFrame(rows[:budget], schema={h: pl.String for h in headers}, orient="row")
        truncated = len(rows) > budget
        notes.append(f"CSV: разделитель {separator!r}, кодировка {encoding}. Пустые записи пропускаются; пустые ячейки сохраняются как строки.")
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet and sheet not in book.sheetnames:
                raise ValueError(f"Лист {sheet!r} не найден.")
            ws = book[sheet] if sheet else book.worksheets[0]
            iterator = ws.iter_rows(values_only=True)
            raw_headers = next(iterator, None)
            if raw_headers is None:
                raise ValueError("Лист пуст.")
            headers = [str(h) if h is not None else "" for h in raw_headers]
            budget = _row_budget(headers, max_rows)
            rows = []
            formulas = False
            for row in iterator:
                formulas |= any(isinstance(v, str) and v.startswith("=") for v in row)
                rows.append([None if v is None else str(v) for v in row])
                if len(rows) > budget:
                    break
            frame = pl.DataFrame(rows[:budget], schema={h: pl.String for h in headers}, orient="row")
            truncated = len(rows) > budget
            notes.append(f"XLSX, лист «{ws.title}»: значения представлены текстом; форматирование ячеек не анализируется.")
            if formulas:
                notes.append("Формулы XLSX показаны как текст, не вычислялись. Для анализа результатов сохраните лист как значения.")
        finally:
            book.close()
    elif ext == ".parquet":
        lazy = pl.scan_parquet(path)
        budget = _row_budget(lazy.collect_schema().names(), max_rows)
        frame = lazy.head(budget + 1).collect()
        truncated = frame.height > budget
        frame = frame.head(budget)
        # Nonfinite numeric values are not conflated with actual nulls.
    else:
        raise ValueError("Поддерживаются CSV, XLSX и Parquet.")
    if frame.width > 500:
        raise ValueError("MVP поддерживает до 500 столбцов.")
    if truncated:
        notes.append(f"Проанализированы только первые {budget:,} записей. Все метрики и кандидаты в ключи относятся к этому фрагменту.")
    if budget < max_rows:
        notes.append(f"Лимит чтения уменьшен до {budget:,} записей: защита от превышения {MAX_CELLS:,} ячеек.")
    unsupported = [name for name, dtype in frame.schema.items() if dtype.is_nested() or dtype == pl.Binary]
    if unsupported:
        raise ValueError(f"В MVP поддерживаются скалярные столбцы. Преобразуйте вложенные/бинарные поля: {', '.join(unsupported)}")
    return Dataset(frame, str(path.resolve()), ext.lstrip("."), truncated, notes)


def _validate_headers(headers: list[str]) -> None:
    if not headers or any(not h.strip() for h in headers):
        raise ValueError("Нужен непустой заголовок для каждого столбца.")
    if len(set(headers)) != len(headers):
        raise ValueError("Заголовки столбцов должны быть уникальными.")
